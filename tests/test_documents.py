"""Round-trip checks for local resume document capabilities."""

from __future__ import annotations

import unittest
from io import BytesIO

from docx import Document
from pypdf import PdfReader, PdfWriter

from job_agent.documents import (
    DocumentError,
    export_resume_docx,
    export_resume_pdf,
    extract_document_text,
)
from job_agent.skills import list_local_document_skills


class DocumentTests(unittest.TestCase):
    SAMPLE = "# 张三\n## 教育经历\n- 示例大学 统计学\n## 技能\n- Python 与 SQL"

    def test_docx_import_includes_paragraphs_and_table_cells(self) -> None:
        document = Document()
        document.add_paragraph("张三的基础简历")
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = "技能"
        row.cells[1].text = "Python"
        output = BytesIO()
        document.save(output)
        text = extract_document_text("resume.docx", output.getvalue())
        self.assertIn("张三的基础简历", text)
        self.assertIn("技能 | Python", text)

    def test_docx_and_pdf_exports_keep_chinese_resume_text(self) -> None:
        docx_data = export_resume_docx(self.SAMPLE)
        word = Document(BytesIO(docx_data))
        self.assertIn("张三", [paragraph.text for paragraph in word.paragraphs])
        self.assertIn("Python 与 SQL", [paragraph.text for paragraph in word.paragraphs])

        pdf_data = export_resume_pdf(self.SAMPLE)
        self.assertTrue(pdf_data.startswith(b"%PDF"))
        text = extract_document_text("resume.pdf", pdf_data)
        self.assertIn("张三", text)
        self.assertIn("Python 与 SQL", text)
        self.assertEqual(len(PdfReader(BytesIO(pdf_data)).pages), 1)

    def test_rejects_corrupt_and_image_only_documents(self) -> None:
        with self.assertRaises(DocumentError):
            extract_document_text("resume.docx", b"not a word file")
        with self.assertRaises(DocumentError):
            extract_document_text("resume.pdf", b"not a pdf")
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        output = BytesIO()
        writer.write(output)
        with self.assertRaisesRegex(DocumentError, "OCR"):
            extract_document_text("scan.pdf", output.getvalue())

    def test_local_skills_are_explicitly_registered(self) -> None:
        self.assertEqual(
            {skill.key for skill in list_local_document_skills()},
            {"document_text_import", "resume_document_export"},
        )
        self.assertTrue(all(not skill.external_data for skill in list_local_document_skills()))


if __name__ == "__main__":
    unittest.main()
