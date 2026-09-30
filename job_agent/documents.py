"""Local DOCX/PDF text import and resume export."""

from __future__ import annotations

import re
import zipfile
from html import escape
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED = 30 * 1024 * 1024
MAX_PDF_PAGES = 30


class DocumentError(ValueError):
    """A file cannot be imported or exported safely."""


def extract_document_text(
    filename: str,
    data: bytes,
    *,
    max_chars: int = 30_000,
) -> str:
    """Read text from a small DOCX or text-based PDF without storing the file."""
    suffix = Path(filename).suffix.lower()
    if suffix not in {".docx", ".pdf"}:
        raise DocumentError("仅支持 .docx 和 .pdf 文件。")
    if not data or len(data) > MAX_FILE_BYTES:
        raise DocumentError("文件为空或超过 8 MB。")
    if suffix == ".docx":
        if not zipfile.is_zipfile(BytesIO(data)):
            raise DocumentError("Word 文件格式不正确。")
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise DocumentError("Word 文件解压后过大。")
        try:
            document = Document(BytesIO(data))
            parts: list[str] = []
            for block in document.iter_inner_content():
                if hasattr(block, "rows"):
                    for row in block.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            parts.append(" | ".join(cells))
                elif block.text.strip():
                    parts.append(block.text.strip())
            text = "\n".join(parts)
        except (ValueError, KeyError, zipfile.BadZipFile) as exc:
            raise DocumentError("无法读取 Word 文件内容。") from exc
    else:
        if not data.startswith(b"%PDF-"):
            raise DocumentError("PDF 文件格式不正确。")
        try:
            reader = PdfReader(BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise DocumentError("PDF 已加密，请先移除密码。")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise DocumentError(f"PDF 超过 {MAX_PDF_PAGES} 页。")
            text = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("无法读取 PDF 文件内容。") from exc
    text = text.strip()
    if not text:
        raise DocumentError("文件没有可提取的文字；扫描版 PDF 需要先进行 OCR。")
    if len(text) > max_chars:
        raise DocumentError(f"提取内容超过 {max_chars:,} 字符，请拆分文件后导入。")
    return text


def _resume_lines(markdown: str) -> list[tuple[str, str]]:
    lines: list[tuple[str, str]] = []
    for raw in markdown.splitlines():
        line = raw.strip()
        if not line or line in {"---", "***"}:
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            kind = "title" if len(heading.group(1)) == 1 else "heading"
            content = heading.group(2)
        elif re.match(r"^[-*]\s+", line):
            kind, content = "bullet", re.sub(r"^[-*]\s+", "", line)
        else:
            kind, content = "body", line
        content = re.sub(r"(?<!\*)\*\*(.+?)\*\*", r"\1", content)
        content = content.replace("`", "").strip()
        if content:
            lines.append((kind, content))
    if not lines:
        raise DocumentError("简历内容为空，无法导出。")
    return lines


def _set_docx_font(style, size: float, *, bold: bool = False) -> None:
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(29, 41, 57)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def export_resume_docx(markdown: str) -> bytes:
    """Make an editable Word resume from reviewed Markdown text."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    _set_docx_font(document.styles["Normal"], 10.5)
    title_style = document.styles.add_style("Resume Name", WD_STYLE_TYPE.PARAGRAPH)
    title_style.base_style = document.styles["Normal"]
    _set_docx_font(title_style, 20, bold=True)
    title_style.paragraph_format.space_after = Pt(12)
    _set_docx_font(document.styles["Heading 2"], 12, bold=True)
    document.styles["Normal"].paragraph_format.space_after = Pt(5)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(5)
    for kind, content in _resume_lines(markdown):
        if kind == "title":
            paragraph = document.add_paragraph(content, style=title_style)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif kind == "heading":
            document.add_paragraph(content, style="Heading 2")
        elif kind == "bullet":
            document.add_paragraph(content, style="List Bullet")
        else:
            document.add_paragraph(content)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_font() -> str:
    name = "AzVResumeChinese"
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    candidates = (
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/Deng.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
    )
    for candidate in candidates:
        if not candidate.exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont(name, str(candidate)))
            return name
        except Exception:
            continue
    raise DocumentError("未找到可嵌入的中文字体，暂时无法导出 PDF。")


def export_resume_pdf(markdown: str) -> bytes:
    """Make a paginated PDF using an embedded Chinese font."""
    font = _pdf_font()
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4,
        leftMargin=48, rightMargin=48, topMargin=42, bottomMargin=42,
        title="定制简历",
    )
    styles = {
        "title": ParagraphStyle(
            "ResumeTitle", fontName=font, fontSize=20, leading=27,
            textColor=colors.HexColor("#1D2939"), alignment=TA_LEFT, spaceAfter=13,
            wordWrap="CJK",
        ),
        "heading": ParagraphStyle(
            "ResumeHeading", fontName=font, fontSize=12, leading=18,
            textColor=colors.HexColor("#1D2939"), spaceBefore=12, spaceAfter=5,
            wordWrap="CJK", keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "ResumeBody", fontName=font, fontSize=10.5, leading=16,
            textColor=colors.HexColor("#263047"), spaceAfter=5,
            wordWrap="CJK",
        ),
        "bullet": ParagraphStyle(
            "ResumeBullet", fontName=font, fontSize=10.5, leading=16,
            textColor=colors.HexColor("#263047"), leftIndent=14,
            firstLineIndent=-10, spaceAfter=5, wordWrap="CJK",
        ),
    }
    story = []
    for kind, content in _resume_lines(markdown):
        prefix = "- " if kind == "bullet" else ""
        story.append(Paragraph(escape(prefix + content), styles[kind]))
        if kind == "heading":
            story.append(Spacer(1, 1))
    document.build(story)
    return output.getvalue()
